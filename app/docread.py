# -*- coding: utf-8 -*-
"""本地文档解析引擎：Word / PDF / Excel / PPT / HTML / CSV / JSON / 文本 / 代码。

全部使用本地库解析（python-docx / pdfminer / PyPDF2 / openpyxl / bs4），不联网。
图片不在这里处理（需要多模态视觉模型），由 agent 的 read_image 工具负责。
"""
import os
import re
import zipfile
import logging

# pdfminer 对某些字体会刷大量无害告警，压到 ERROR 级，保持控制台干净
logging.getLogger("pdfminer").setLevel(logging.ERROR)

MAX_CHARS = 30000

DOC_EXTS = {".docx", ".docm", ".dotx"}
PDF_EXTS = {".pdf"}
XLS_EXTS = {".xlsx", ".xlsm", ".xltx"}
PPT_EXTS = {".pptx", ".pptm", ".potx"}
HTML_EXTS = {".html", ".htm", ".xhtml"}
TEXT_EXTS = {".txt", ".md", ".markdown", ".log", ".json", ".csv", ".tsv",
             ".ini", ".cfg", ".conf", ".yaml", ".yml", ".xml", ".rtf", ".srt",
             ".py", ".js", ".ts", ".jsx", ".tsx", ".java", ".c", ".cpp", ".cc",
             ".h", ".hpp", ".cs", ".go", ".rs", ".php", ".rb", ".sh", ".bat",
             ".ps1", ".sql", ".css", ".vue", ".kt", ".swift"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".bmp", ".gif", ".webp", ".tif", ".tiff"}

_KIND = {
    ".docx": "Word 文档", ".docm": "Word 文档", ".dotx": "Word 模板", ".doc": "Word 文档(旧)",
    ".pdf": "PDF 文档",
    ".xlsx": "Excel 表格", ".xlsm": "Excel 表格", ".xltx": "Excel 模板", ".xls": "Excel(旧)",
    ".pptx": "PPT 演示文稿", ".pptm": "PPT 演示文稿", ".potx": "PPT 模板",
    ".html": "网页", ".htm": "网页", ".xhtml": "网页",
    ".csv": "CSV 表格", ".tsv": "TSV 表格", ".json": "JSON 数据",
    ".md": "Markdown 文档", ".txt": "文本文件", ".log": "日志文件",
    ".png": "图片", ".jpg": "图片", ".jpeg": "图片", ".bmp": "图片",
    ".webp": "图片", ".gif": "图片", ".tif": "图片", ".tiff": "图片",
}


def ext_of(path):
    return os.path.splitext(path or "")[1].lower()


def file_kind(path):
    """返回人类可读的文件类型，如 'Word 文档'。"""
    return _KIND.get(ext_of(path), "文件")


def is_image(path):
    return ext_of(path) in IMAGE_EXTS


def is_supported(path):
    e = ext_of(path)
    return e in (DOC_EXTS | PDF_EXTS | XLS_EXTS | PPT_EXTS
                 | HTML_EXTS | TEXT_EXTS | IMAGE_EXTS)


def _clip(text, max_chars):
    text = text or ""
    if max_chars and len(text) > max_chars:
        return text[:max_chars] + f"\n\n...（内容过长，已截断；原文约 {len(text)} 字）"
    return text


def _read_text(path, max_chars):
    """按常见中文编码依次尝试读取纯文本。"""
    last = None
    for enc in ("utf-8-sig", "utf-8", "gbk", "cp936", "latin-1"):
        try:
            with open(path, "r", encoding=enc) as f:
                return f.read()
        except Exception as e:
            last = e
            continue
    # 实在不行按二进制容错
    with open(path, "rb") as f:
        return f.read().decode("utf-8", errors="replace")


def _read_docx(path, max_chars):
    parts = []
    try:
        from docx import Document
        doc = Document(path)
        for p in doc.paragraphs:
            t = (p.text or "").strip()
            if t:
                parts.append(t)
        for i, tbl in enumerate(doc.tables):
            parts.append(f"\n[表格 {i + 1}]")
            for row in tbl.rows:
                cells = [(c.text or "").strip().replace("\n", " ") for c in row.cells]
                parts.append(" | ".join(cells))
        return "\n".join(parts)
    except ImportError:
        pass
    # 兜底：直接解 docx（zip）里的 document.xml
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf-8", "ignore")
    xml = xml.replace("</w:p>", "</w:p>\n")
    texts = re.findall(r"<w:t[^>]*>(.*?)</w:t>", xml, re.DOTALL)
    return "\n".join(t for t in texts if t.strip())


def _read_pdf(path, max_chars):
    txt = ""
    try:
        from pdfminer.high_level import extract_text
        try:
            txt = extract_text(path, maxpages=60)
        except TypeError:
            txt = extract_text(path)
    except Exception:
        txt = ""
    if not (txt or "").strip():
        try:
            from PyPDF2 import PdfReader
            reader = PdfReader(path)
            pages = []
            for i, pg in enumerate(reader.pages):
                if i >= 80:
                    break
                try:
                    pages.append(pg.extract_text() or "")
                except Exception:
                    pages.append("")
            txt = "\n".join(pages)
        except Exception as e:
            raise RuntimeError(f"PDF 解析失败：{e}")
    return txt


def _read_xlsx(path, max_chars):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    parts = []
    try:
        for ws in wb.worksheets:
            parts.append(f"\n[工作表] {ws.title}")
            n = 0
            for row in ws.iter_rows(values_only=True):
                cells = [("" if c is None else str(c)) for c in row]
                if any(x.strip() for x in cells):
                    parts.append(" | ".join(cells))
                    n += 1
                if n >= 300:
                    parts.append("...（本表仅显示前 300 行）")
                    break
    finally:
        try:
            wb.close()
        except Exception:
            pass
    return "\n".join(parts)


def _read_pptx(path, max_chars):
    try:
        from pptx import Presentation
        prs = Presentation(path)
        parts = []
        for i, slide in enumerate(prs.slides):
            parts.append(f"\n[幻灯片 {i + 1}]")
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    txt = (shape.text_frame.text or "").strip()
                    if txt:
                        parts.append(txt)
        return "\n".join(parts)
    except Exception:
        pass
    # 兜底：解 pptx（zip）里的 slideN.xml
    parts = []
    with zipfile.ZipFile(path) as z:
        names = [n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)]
        names.sort(key=lambda s: int(re.findall(r"\d+", s)[0]))
        for n in names:
            xml = z.read(n).decode("utf-8", "ignore")
            texts = re.findall(r"<a:t>(.*?)</a:t>", xml, re.DOTALL)
            parts.append(f"\n[{n}] " + " ".join(texts))
    return "\n".join(parts)


def _read_html(path, max_chars):
    raw = _read_text(path, None)
    try:
        from bs4 import BeautifulSoup
        try:
            soup = BeautifulSoup(raw, "lxml")
        except Exception:
            soup = BeautifulSoup(raw, "html.parser")
        for tag in soup(["script", "style", "noscript"]):
            tag.decompose()
        txt = soup.get_text("\n")
    except Exception:
        txt = re.sub(r"<[^>]+>", " ", raw)
    txt = re.sub(r"[ \t]+", " ", txt)
    txt = re.sub(r"\n{3,}", "\n\n", txt)
    return txt.strip()


def read_document(path, max_chars=MAX_CHARS):
    """解析文档并返回纯文本。失败抛异常（附中文原因）。"""
    if not path:
        raise ValueError("未提供文件路径")
    if not os.path.isfile(path):
        raise FileNotFoundError(f"文件不存在：{path}")
    e = ext_of(path)
    size = os.path.getsize(path)

    if e in DOC_EXTS:
        txt = _read_docx(path, max_chars)
    elif e in PDF_EXTS:
        if size > 60 * 1024 * 1024:
            raise ValueError("PDF 超过 60MB，暂不解析")
        txt = _read_pdf(path, max_chars)
    elif e in XLS_EXTS:
        txt = _read_xlsx(path, max_chars)
    elif e in PPT_EXTS:
        txt = _read_pptx(path, max_chars)
    elif e in HTML_EXTS:
        txt = _read_html(path, max_chars)
    elif e in TEXT_EXTS or not e:
        txt = _read_text(path, max_chars)
    elif e == ".doc":
        raise ValueError("旧版 .doc 无法直接解析，请先用 Word 另存为 .docx")
    elif e == ".xls":
        raise ValueError("旧版 .xls 无法直接解析，请先用 Excel 另存为 .xlsx")
    else:
        # 未知类型：先按文本试读，失败则明确报错
        try:
            txt = _read_text(path, max_chars)
        except Exception:
            raise ValueError(f"暂不支持解析该类型：{e or '未知'}")

    txt = _clip(txt, max_chars)
    if not (txt or "").strip():
        txt = ("（已成功解析文件，但没有提取到文字——可能是扫描件/图片型文档；"
               "若是这种情况，请把该页截图后用图片识别。）")
    return txt


def describe(path):
    """给 UI 用：返回 (类型标签, 大小可读串)。"""
    if not os.path.exists(path):
        return ("不存在", "")
    try:
        size = os.path.getsize(path)
    except Exception:
        size = 0
    for unit in ("B", "KB", "MB", "GB"):
        if size < 1024 or unit == "GB":
            human = f"{size:.0f} {unit}" if unit == "B" else f"{size:.1f} {unit}"
            break
        size /= 1024
    return (file_kind(path), human)
