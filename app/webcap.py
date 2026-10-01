# -*- coding: utf-8 -*-
"""网页采集能力（v9.15.0）：带会话抓取 / 批量表单查询 / GitHub。

给 AI 工作台补上「像人一样操作网页」的能力，解决「查询类网站批量查数据」这类
以前做不了、或者要靠模型自己写爬虫（弱模型必跪）的任务：

  · ``web_scrape``     —— 带 Cookie 会话的 GET/POST；能拿**原始 HTML**（含 script）；
                          能自动列出页面上的**表单结构**（字段名/隐藏 token）；
                          能按表格/字段提取结构化数据。
  · ``web_form_batch`` —— ★ 批量表单查询：一个网址 + 一串值 -> 逐个提交 -> 汇总成表，
                          并**直接落盘 Excel/CSV**（易查分、成绩查询、考勤/分班查询
                          这类站点的通用解法，支持断点续查）。
  · ``github``         —— 读仓库文件/列目录/搜索/看 issue 与 PR；优先用本机已登录的
                          ``gh`` CLI，装不了就回落 GitHub 公开 HTTP 接口。

设计约束：
- HTTP 一律走 ``requests``（会自动走系统代理）；本机 urllib 会被中间代理拦，只在不
  得不已时才用它兜底。
- **会话池**按名字复用，让「先访问页面拿 Cookie、再提交表单」这种两段式自动串起来。
- 只做正常读取与表单提交：不爆破、不绕验证码。**一旦遇到验证码/风控就立刻停下并
  如实报告**（绝不假装查到了）。
"""
import csv
import io
import json
import os
import re
import subprocess
import time
from urllib.parse import urljoin, urlparse

# --------------------------------------------------------------------------
# 会话池：让多步操作（预热 -> 提交 -> 取结果）共用同一份 Cookie
# --------------------------------------------------------------------------
_SESSIONS = {}

DEFAULT_UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
              "(KHTML, like Gecko) Chrome/126.0.0.0 Safari/537.36")

# 结果页里这些噪声行不要
_NOISE = re.compile(
    r"^(查询结果|查询必读|分享好友|订阅通知|我的保存|本查询使用|易查分|"
    r"免费制作查询|对查询结果有疑问请联系我们|长按关注查询公众号|"
    r"快速进入查询|保存结果|接收通知|发成绩用|请输入正确的查询条件进行查询|"
    r"查看说明|关闭|提示|查询中\.\.\.)$")


def _clean(s):
    s = re.sub(r"[\u200b\ufeff\xa0]", " ", str(s or ""))
    s = re.sub(r"\s+", " ", s).strip()
    return s


def _session(name=None, headers=None, cookies=None, url=None):
    """取（或建）一个会话。name 相同 -> 复用同一份 Cookie。"""
    key = str(name or "").strip()
    if not key and url:
        try:
            key = urlparse(url).netloc
        except Exception:
            key = ""
    key = key or "default"
    s = _SESSIONS.get(key)
    if s is None:
        import requests
        s = requests.Session()
        s.headers.update({
            "User-Agent": DEFAULT_UA,
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
            "Accept-Language": "zh-CN,zh;q=0.9",
        })
        _SESSIONS[key] = s
    if isinstance(headers, dict):
        for k, v in headers.items():
            if v is not None:
                s.headers[str(k)] = str(v)
    if isinstance(cookies, dict):
        for k, v in cookies.items():
            s.cookies.set(str(k), str(v))
    if isinstance(cookies, str) and cookies.strip():
        for part in cookies.split(";"):
            if "=" in part:
                k, v = part.split("=", 1)
                s.cookies.set(k.strip(), v.strip())
    return s


def reset_session(name=None):
    """丢掉会话（比如要重新登录/换账号）。"""
    if name is None:
        _SESSIONS.clear()
        return
    _SESSIONS.pop(str(name), None)


def _pick_encoding(resp):
    enc = (resp.encoding or "").lower()
    if enc and enc not in ("iso-8859-1", "ascii"):
        return resp.encoding
    ctype = (resp.headers.get("content-type") or "").lower()
    m = re.search(r"charset=([\w\-]+)", ctype)
    if m:
        return m.group(1)
    # 从 meta 里找
    head = resp.content[:3000].decode("ascii", "ignore")
    m = re.search(r'charset=["\']?([\w\-]+)', head, re.I)
    if m:
        return m.group(1)
    return "utf-8"


def _text_of(resp, limit=None):
    enc = _pick_encoding(resp)
    try:
        t = resp.content.decode(enc, "replace")
    except Exception:
        t = resp.content.decode("utf-8", "replace")
    if limit and len(t) > limit:
        t = t[:limit] + f"\n…（共 {len(t)} 字符，已截断）"
    return t


def _soup(html):
    from bs4 import BeautifulSoup
    try:
        return BeautifulSoup(html, "lxml")
    except Exception:
        return BeautifulSoup(html, "html.parser")


def page_text(html):
    """正文纯文本（去 script/style，去噪声行）。"""
    soup = _soup(html)
    for t in soup(["script", "style", "noscript"]):
        t.decompose()
    raw = soup.get_text("\n")
    lines = []
    for ln in raw.split("\n"):
        ln = _clean(ln)
        if not ln or _NOISE.match(ln):
            continue
        if lines and lines[-1] == ln:
            continue
        lines.append(ln)
    return "\n".join(lines)


def extract_pairs(html):
    """从结果页抠出「字段 -> 值」对。

    优先易查分的结果表（``js_result_table``），其次 weui 单元格，再退到通用两列表格。
    这是「输入姓名 -> 返回一张信息表」这类查询站点的通用解法。
    """
    soup = _soup(html)
    pairs = []

    def add(k, v):
        k, v = _clean(k), _clean(v)
        if not k and not v:
            return
        if k and _NOISE.match(k):
            return
        pairs.append([k, v])

    # ① 易查分结果表
    for tb in soup.select("table.js_result_table, table.s_table-bordered, "
                          "#result_data_table table"):
        for tr in tb.find_all("tr"):
            tds = tr.find_all(["td", "th"])
            if len(tds) >= 2:
                add(tds[0].get_text(" "), tds[1].get_text(" "))
        if pairs:
            return pairs

    # ② weui 表单式单元格
    for cell in soup.select(".weui-cell"):
        lab = cell.select_one(".weui-label")
        bd = cell.select_one(".weui-cell__bd")
        if lab and bd:
            add(lab.get_text(" "), bd.get_text(" "))
    if pairs:
        return pairs

    # ③ 通用「两列」表格
    for tb in soup.find_all("table"):
        got = []
        for tr in tb.find_all("tr"):
            tds = tr.find_all(["td", "th"])
            if len(tds) == 2:
                k, v = _clean(tds[0].get_text(" ")), _clean(tds[1].get_text(" "))
                if k and len(k) <= 24 and k not in [x[0] for x in got]:
                    got.append([k, v])
        if len(got) >= 2:
            pairs = got
            break
    if pairs:
        return pairs

    # ④ dl / dt+dd
    for dt in soup.find_all("dt"):
        dd = dt.find_next_sibling("dd")
        if dd:
            add(dt.get_text(" "), dd.get_text(" "))
    return pairs


def form_list(html):
    """列出页面上的表单结构（字段名 / 类型 / 隐藏值 / action）——排查接口时最有用。"""
    soup = _soup(html)
    forms = soup.find_all("form")
    out = []
    for i, f in enumerate(forms):
        rec = {
            "index": i,
            "action": f.get("action"),
            "method": (f.get("method") or "GET").upper(),
            "id": f.get("id"),
            "fields": [],
        }
        for el in f.find_all(["input", "select", "textarea", "button"]):
            rec["fields"].append({
                "tag": el.name,
                "name": el.get("name"),
                "id": el.get("id"),
                "type": el.get("type"),
                "value": (el.get("value") or "")[:80] or None,
            })
        out.append(rec)
    if not out:
        # 表单可能是 JS 拼的：至少把裸 input 列出来
        loose = [{"tag": el.name, "name": el.get("name"), "id": el.get("id"),
                  "type": el.get("type"),
                  "value": (el.get("value") or "")[:80] or None}
                 for el in soup.find_all(["input", "select", "textarea"])]
        if loose:
            out.append({"index": -1, "action": None, "method": "?",
                        "id": "（无 form 标签，字段为全页 input）",
                        "fields": loose})
    # ★ 关键补充：提交地址常常藏在 JS 里（页面 form 没有 action）。
    eps = sniff_endpoints(html)
    if eps:
        out.append({
            "index": -2, "action": None, "method": "(JS 里的提交地址)",
            "id": "★ 从页面 JavaScript 里嗅探到的接口（表单没有 action 时看这里）",
            "fields": [{"tag": "endpoint", "name": e, "id": None,
                        "type": "url", "value": None} for e in eps],
        })
    return out


_ENDPOINT_HINT = re.compile(
    r"(verifycondition|submit|search|query|check|result|login|ajax|api|action)",
    re.I)
_ENDPOINT_PATS = [
    r"""\$\.post\(\s*["']([^"']+)["']""",
    r"""\$\.get\(\s*["']([^"']+)["']""",
    r"""\$\.ajax\(\s*\{[^}]{0,400}?url\s*:\s*["']([^"']+)["']""",
    r"""url\s*:\s*["']([^"']{4,200})["']""",
    r"""fetch\(\s*["']([^"']+)["']""",
    r"""axios\.(?:post|get)\(\s*["']([^"']+)["']""",
    r"""\.open\(\s*["'](?:POST|GET)["']\s*,\s*["']([^"']+)["']""",
]


def sniff_endpoints(html, base=None):
    """从页面 HTML/JS 里嗅探「表单提交接口」。

    很多查询类网站（易查分就是）的 form 没有 action，真正的提交地址写在
    ``$.post("...")`` 里。模型只能拿到页面 URL，所以这一步必须程序自己做。
    按「像接口地址」的程度排序，靠谱的排前面。
    """
    if not html:
        return []
    found = []
    for pat in _ENDPOINT_PATS:
        for m in re.finditer(pat, html, re.I):
            u = _clean(m.group(1))
            if not u or u.startswith(("javascript:", "#", "data:")):
                continue
            if len(u) < 4 or len(u) > 300:
                continue
            if u not in found:
                found.append(u)
    def score(u):
        s = 0
        if _ENDPOINT_HINT.search(u):
            s -= 10
        if u.endswith(".html") or u.endswith(".php") or u.endswith(".do"):
            s -= 3
        if "{{" in u or "${" in u:
            s += 20           # 模板占位符，基本没用
        if u.startswith("/"):
            s -= 2
        return s
    found.sort(key=score)
    out = []
    for u in found[:6]:
        if base:
            try:
                u = urljoin(base, u)
            except Exception:
                pass
        out.append(u)
    return out


def extract_tables(html):
    """把页面里的表格转成 Markdown 表格。"""
    soup = _soup(html)
    chunks = []
    for i, tb in enumerate(soup.find_all("table")):
        rows = []
        for tr in tb.find_all("tr"):
            cells = [_clean(td.get_text(" ")) for td in tr.find_all(["td", "th"])]
            if any(cells):
                rows.append(cells)
        if len(rows) < 1:
            continue
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]
        md = ["| " + " | ".join(rows[0]) + " |",
              "|" + "---|" * width]
        for r in rows[1:]:
            md.append("| " + " | ".join(r) + " |")
        chunks.append(f"[表 {i + 1}]\n" + "\n".join(md))
    return "\n\n".join(chunks)


def _load_values(values, values_from):
    """把「一串值」凑齐：可以直接给列表，也可以从 Excel/CSV/txt 第一列读。"""
    vals = []
    if isinstance(values, str):
        values = [x for x in re.split(r"[\n,，、;；]+", values)]
    if isinstance(values, (list, tuple)):
        vals.extend(_clean(v) for v in values if _clean(v))
    if values_from:
        p = str(values_from)
        ext = os.path.splitext(p)[1].lower()
        if ext in (".xlsx", ".xlsm", ".xls"):
            try:
                import openpyxl
                wb = openpyxl.load_workbook(p, read_only=True, data_only=True)
                for ws in wb.worksheets:
                    for row in ws.iter_rows(values_only=True):
                        for cell in row:
                            if cell is not None and _clean(cell):
                                vals.append(_clean(cell))
                                break
                    break
                wb.close()
            except Exception as e:
                raise RuntimeError(f"读 Excel 失败：{e}")
        else:
            raw = None
            for enc in ("utf-8-sig", "utf-8", "gbk"):
                try:
                    raw = io.open(p, encoding=enc).read()
                    break
                except Exception:
                    continue
            if raw is None:
                raise RuntimeError(f"读文件失败：{p}")
            for ln in raw.split("\n"):
                ln = _clean(ln)
                if not ln:
                    continue
                vals.append(ln.split(",")[0].strip().strip('"'))
    # 去重保序
    seen, out = set(), []
    for v in vals:
        if v and v not in seen:
            seen.add(v)
            out.append(v)
    return out


# --------------------------------------------------------------------------
# 工具一：web_scrape
# --------------------------------------------------------------------------
def scrape(url, method="GET", form=None, json_body=None, headers=None,
           cookies=None, session=None, extract="auto", selector=None,
           raw=False, max_chars=9000, timeout=30, referer=None):
    """带会话抓取一个网址，按需返回正文/原始HTML/表单结构/表格。"""
    import requests

    url = str(url or "").strip()
    if not url.lower().startswith(("http://", "https://")):
        return "[错误] url 必须以 http:// 或 https:// 开头"
    method = str(method or "GET").upper()
    s = _session(session, headers, cookies, url)
    hd = {}
    if referer:
        hd["Referer"] = referer

    kw = {"timeout": int(timeout or 30), "allow_redirects": True}
    if hd:
        kw["headers"] = hd
    try:
        if method == "GET":
            resp = s.get(url, **kw)
        else:
            data, js = None, None
            if isinstance(form, dict):
                data = {str(k): ("" if v is None else str(v))
                        for k, v in form.items()}
            elif isinstance(form, str) and form.strip():
                data = form
            if json_body is not None:
                js = json_body
            if data is None and js is None:
                data = {}
            kw["headers"] = dict(hd, **{
                "X-Requested-With": "XMLHttpRequest",
                "Origin": f"{urlparse(resp_url(url)).scheme}://"
                          f"{urlparse(resp_url(url)).netloc}",
                "Referer": referer or url,
            })
            resp = s.post(url, data=data, json=js, **kw)
    except Exception as e:
        return f"[错误] 请求失败：{type(e).__name__}: {e}"

    ctype = (resp.headers.get("content-type") or "").lower()
    body = _text_of(resp, limit=None)

    head = (f"[{'成功' if 200 <= resp.status_code < 400 else 'HTTP ' + str(resp.status_code)}] "
            f"{method} {url}\n状态码：{resp.status_code}  类型：{ctype or '?'}  "
            f"长度：{len(resp.content)} 字节")

    mode = str(extract or "auto").lower()
    if raw or mode == "html":
        cut = body[:max_chars]
        tail = "" if len(body) <= max_chars else f"\n…（共 {len(body)} 字符，已截断）"
        return head + "\n--- 原始 HTML ---\n" + cut + tail

    if mode == "forms" or (mode == "auto" and "form" in body.lower()
                           and "type=\"text\"" in body.lower() and not raw):
        fl = form_list(body)
        if fl:
            return head + "\n--- 页面表单结构 ---\n" + json.dumps(
                fl, ensure_ascii=False, indent=2)

    if mode == "json" or (mode == "auto" and "json" in ctype):
        try:
            return head + "\n--- JSON ---\n" + json.dumps(
                resp.json(), ensure_ascii=False, indent=2)[:max_chars]
        except Exception:
            pass

    if mode == "tables":
        tb = extract_tables(body)
        if tb:
            return head + "\n--- 表格 ---\n" + tb[:max_chars]

    if mode == "fields":
        pr = extract_pairs(body)
        if pr:
            return head + "\n--- 字段 ---\n" + "\n".join(
                f"{k}：{v}" for k, v in pr)

    if selector:
        try:
            nodes = _soup(body).select(str(selector))
            got = "\n".join(_clean(n.get_text(" ")) for n in nodes if _clean(n.get_text(" ")))
            return head + f"\n--- 选择器 {selector} ---\n" + (got or "（没匹配到内容）")
        except Exception as e:
            return head + f"\n[错误] 选择器解析失败：{e}"

    # auto：先试字段对（查询结果页常见），再给正文
    txt = page_text(body)
    if mode == "auto":
        pr = extract_pairs(body)
        if len(pr) >= 2:
            txt = "\n".join(f"{k}：{v}" for k, v in pr) + "\n\n（原始正文）\n" + txt
    if len(txt) > max_chars:
        txt = txt[:max_chars] + f"\n…（共 {len(txt)} 字符，已截断）"
    return head + "\n--- 正文 ---\n" + txt


def resp_url(u):
    return u


# --------------------------------------------------------------------------
# 工具二：web_form_batch —— 批量表单查询
# --------------------------------------------------------------------------
def form_batch(url, field, values=None, values_from=None, extra=None,
               method="POST", session=None, delay=1.5, max_items=None,
               save_path=None, timeout=30, resume=True, headers=None,
               cookies=None, label="查询值", stop_on_block=True,
               fetch_result=True, endpoint=None):
    """批量往一个表单里灌不同的值，收集每次的结果。

    专治「输入姓名 -> 返回我的宿舍/成绩/录取信息」这类查询网站：
      ① 每个值提交一次（默认 POST 表单）；
      ② 若返回值是一个 JSON，且里面有 url/redirect 之类字段 -> 自动跟进去拿结果页；
      ③ 从结果页抠出「字段：值」，汇总成表；
      ④ 可选直接把整张表写成 Excel/CSV（这就是交付产物）。

    返回 ``(给模型看的文本, 落盘路径或 None, 落盘字节数)``。
    """
    import requests

    url = str(url or "").strip()
    if not url.lower().startswith(("http://", "https://")):
        return "[错误] url 必须以 http:// 或 https:// 开头", None, 0
    fld = str(field or "").strip()
    if not fld:
        return "[错误] 需要 field 参数：表单里那个字段的名字（如姓名站点的 s_xingming）", None, 0

    try:
        vals = _load_values(values, values_from)
    except Exception as e:
        return f"[错误] {e}", None, 0
    if not vals:
        return ("[错误] 没有拿到要查询的值。请用 values 直接给一串，"
                "或用 values_from 指向 Excel/CSV/txt 文件。"), None, 0

    limit = int(max_items) if max_items else len(vals)
    vals = vals[:limit]

    # ---- 断点续查：把已经查到过的值跳过 ----
    done = {}
    if resume and save_path and os.path.isfile(save_path):
        try:
            done = _read_done(save_path)
        except Exception:
            done = {}

    s = _session(session, headers, cookies, url)
    origin = f"{urlparse(url).scheme}://{urlparse(url).netloc}"
    hd = {"X-Requested-With": "XMLHttpRequest", "Origin": origin,
          "Referer": url}

    # 预热一次：既让站点先发 Cookie（很多查询站靠 PHPSESSID 串状态），
    # 也顺便**嗅探真正的提交接口** —— 给出的 url 往往是「查询页面」，
    # 而真正的提交地址藏在页面 JS 的 $.post("...") 里（易查分就是这样）。
    target = str(endpoint or "").strip() or url
    sniffed = None
    if not endpoint:
        try:
            r0 = s.get(url, timeout=int(timeout or 30))
            eps = sniff_endpoints(_text_of(r0), url)
            if eps:
                sniffed = eps[0]
                target = eps[0]
        except Exception:
            pass

    rows = list(done.get("_rows") or [])
    keys = list(done.get("_keys") or [])
    fails = []
    skipped = 0
    blocked = None

    for i, v in enumerate(vals, 1):
        if v in (done.get("_ok") or set()):
            skipped += 1
            continue
        payload = {}
        if isinstance(extra, (dict, str)):
            if isinstance(extra, str):
                try:
                    extra = json.loads(extra)
                except Exception:
                    extra = {}
            if isinstance(extra, dict):
                payload.update({str(k): str(vv) for k, vv in extra.items()})
        payload[fld] = v

        try:
            if method.upper() == "GET":
                resp = s.get(target, params=payload, timeout=int(timeout or 30),
                             headers=hd)
            else:
                resp = s.post(target, data=payload, timeout=int(timeout or 30),
                              headers=hd)
        except Exception as e:
            fails.append(f"{v}: 请求失败 {type(e).__name__}: {e}")
            continue

        html = _text_of(resp)
        j = None
        if "json" in (resp.headers.get("content-type") or "").lower():
            try:
                j = resp.json()
            except Exception:
                j = None

        # 风控 / 失败判定
        if isinstance(j, dict):
            st = j.get("status")
            if j.get("captchaRequired") or j.get("showPicVerify"):
                blocked = (f"网站要求**验证码**（{j.get('info') or '滑块验证'}）。"
                           f"已停在「{v}」这一条，前面的结果都已保存。")
                break
            if st in (0, "0", False) and (j.get("info") or j.get("msg")):
                fails.append(f"{v}: {_clean(j.get('info') or j.get('msg'))}")
                continue
            if fetch_result:
                nxt = _find_result_url(j)
                if nxt:
                    try:
                        r2 = s.get(urljoin(target, nxt), timeout=int(timeout or 30),
                                   headers={"Referer": url})
                        html = _text_of(r2)
                    except Exception:
                        pass

        pr = extract_pairs(html)
        rec = {label: v}
        if pr:
            for k, vv in pr:
                if k and k not in rec:
                    rec[k] = vv
                if k and k not in keys:
                    keys.append(k)
            # 若一个字段都没识别出来，退化成整段正文
            if len(rec) <= 1:
                rec["结果"] = page_text(html)[:300]
                if "结果" not in keys:
                    keys.append("结果")
        else:
            rec["结果"] = page_text(html)[:300] or "（结果页没有可读文字）"
            if "结果" not in keys:
                keys.append("结果")

        rows.append(rec)
        if delay:
            time.sleep(float(delay))

    # ---- 落盘 ----
    saved, size = None, 0
    if save_path:
        try:
            saved, size = _save_rows(save_path, [label] + keys, rows, label)
        except Exception as e:
            fails.append(f"写文件失败：{e}")

    ok_vals = [r.get(label) for r in rows]
    lines = [f"[完成] 批量查询 {len(vals)} 个值"
             f"（本轮实际提交 {len(vals) - skipped - len(fails) - (1 if blocked else 0)} 个"
             f"，跳过已查 {skipped} 个，失败 {len(fails)} 个）"
             if not blocked else
             f"[中断] 只完成了一部分（网站触发了验证码/风控）"]
    lines.append(f"提交地址：{target}")
    if sniffed:
        lines.append(f"（该地址是从页面 JS 里自动嗅探到的，表单本身没有 action）")
    if saved:
        lines.append(f"结果已保存：{saved}（{size} 字节）")
    else:
        lines.append("⚠️ 没有落盘文件（未指定 save_path）")

    # 结果表（给模型看前 60 行）
    cols = [label] + keys
    lines.append("")
    lines.append("| " + " | ".join(cols) + " |")
    lines.append("|" + "---|" * len(cols))
    for r in rows[:60]:
        lines.append("| " + " | ".join(
            str(r.get(c, "")).replace("|", "/") for c in cols) + " |")
    if len(rows) > 60:
        lines.append(f"…（共 {len(rows)} 行，完整内容在落盘文件里）")
    if blocked:
        lines.append("")
        lines.append("⚠️ " + blocked)
    if fails:
        lines.append("")
        lines.append("失败明细（前 15 条）：")
        lines.extend("  · " + f for f in fails[:15])

    text = "\n".join(lines)
    if len(text) > 12000:
        text = text[:12000] + "\n…（已截断，完整内容见落盘文件）"
    _ = (origin, ok_vals)
    return text, saved, size


def _find_result_url(j):
    """从返回 JSON 里找出「结果页地址」字段。"""
    if not isinstance(j, dict):
        return None
    for k in ("url", "Url", "URL", "redirect", "redirectUrl", "result_url",
              "resultUrl", "jumpUrl", "location", "link"):
        v = j.get(k)
        if isinstance(v, str) and v.strip() and v.strip() not in ("", "#"):
            return v.strip()
    for v in j.values():
        if isinstance(v, dict):
            r = _find_result_url(v)
            if r:
                return r
    return None


def _read_done(path):
    """读回上次的结果（支持 csv / xlsx），用于断点续查。"""
    rows, ok = [], set()
    ext = os.path.splitext(path)[1].lower()
    keys = []
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        header = next(it, None)
        if header:
            keys = [str(h) for h in header[1:] if h]
        for r in it:
            if not r or r[0] is None:
                continue
            rec = {str(header[0]): str(r[0])}
            for idx, h in enumerate(header[1:], 1):
                if h and idx < len(r) and r[idx] is not None:
                    rec[str(h)] = str(r[idx])
            rows.append(rec)
            ok.add(str(r[0]))
        wb.close()
    else:
        for enc in ("utf-8-sig", "utf-8", "gbk"):
            try:
                with io.open(path, encoding=enc, newline="") as f:
                    rd = list(csv.reader(f))
                break
            except Exception:
                rd = []
        if rd:
            header = rd[0]
            keys = [h for h in header[1:] if h]
            for r in rd[1:]:
                if not r:
                    continue
                rec = {header[0]: r[0]}
                for i2, h in enumerate(header[1:], 1):
                    if i2 < len(r):
                        rec[h] = r[i2]
                rows.append(rec)
                ok.add(r[0])
    return {"_rows": rows, "_ok": ok, "_keys": keys}


def _save_rows(path, cols, rows, label):
    """把汇总表写成 csv 或 xlsx。返回 (路径, 字节数)。"""
    d = os.path.dirname(os.path.abspath(path))
    if d:
        os.makedirs(d, exist_ok=True)
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "查询结果"
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="4472C4")
            c.alignment = Alignment(horizontal="center")
        for r in rows:
            ws.append([r.get(c, "") for c in cols])
        for i, c in enumerate(cols, 1):
            width = max(10, min(36, max(
                [len(str(c))] + [len(str(r.get(c, ""))) for r in rows[:200]] or [10]) + 4))
            ws.column_dimensions[
                openpyxl.utils.get_column_letter(i)].width = width
        ws.freeze_panes = "A2"
        wb.save(path)
    else:
        if ext not in (".csv", ".txt", ""):
            path = os.path.splitext(path)[0] + ".csv"
        with io.open(path, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(cols)
            for r in rows:
                w.writerow([r.get(c, "") for c in cols])
    return path, os.path.getsize(path)


# --------------------------------------------------------------------------
# 工具三：github
# --------------------------------------------------------------------------
_GH_CANDIDATES = [
    "gh", "gh.exe",
    r"C:\Program Files\GitHub CLI\gh.exe",
    r"C:\Program Files (x86)\GitHub CLI\gh.exe",
    os.path.expandvars(r"%LOCALAPPDATA%\Programs\GitHub CLI\gh.exe"),
]


def _find_gh():
    import shutil as _sh
    for c in _GH_CANDIDATES:
        p = _sh.which(c) if os.sep not in c else (c if os.path.isfile(c) else None)
        if p:
            return p
    return None


def _gh(args, timeout=60):
    exe = _find_gh()
    if not exe:
        return None, "gh CLI 没找到"
    try:
        import sys
        cf = 0
        if sys.platform == "win32":
            cf = 0x08000000  # CREATE_NO_WINDOW
        p = subprocess.run([exe] + args, capture_output=True, timeout=timeout,
                           creationflags=cf)
        out = p.stdout.decode("utf-8", "replace")
        err = p.stderr.decode("utf-8", "replace")
        if p.returncode != 0:
            return None, (err or out or f"gh 退出码 {p.returncode}").strip()[:500]
        return out, None
    except Exception as e:
        return None, f"{type(e).__name__}: {e}"


def _gh_ready():
    out, err = _gh(["auth", "status"])
    if err:
        return False, err
    return True, (out or "").strip()


def github(action="read", repo=None, path=None, ref=None, query=None,
           kind="repos", number=None, state="open", limit=20, token=None,
           timeout=45):
    """GitHub 操作。action: read / list / repo / search / issues / pr / releases。"""
    act = str(action or "read").strip().lower()
    if act in ("file", "get", "cat", "show"):
        act = "read"
    if act in ("dir", "ls", "tree"):
        act = "list"
    if act in ("issue", "issue_list", "prs"):
        act = "issues" if act != "prs" else "pr"
    if act in ("stars", "trending"):
        act = "search"

    repo = (str(repo).strip() if repo else "") or None
    if repo:
        repo = repo.replace("https://github.com/", "").strip("/")
        if repo.endswith(".git"):
            repo = repo[:-4]
    ref = (str(ref).strip() if ref else "") or None
    path = (str(path).strip().lstrip("/") if path else "") or None

    if act in ("search",):
        if not query:
            return "[错误] search 需要 query 参数"
        kind = (str(kind or "repos").lower())
        kmap = {"repo": "repos", "repositories": "repos", "code": "code",
                "issues": "issues", "pr": "prs", "commits": "commits",
                "users": "users"}
        k = kmap.get(kind, "repos")
        out, err = _gh(["search", k, str(query), "--limit", str(int(limit or 20)),
                        "--json", "fullName,description,stargazersCount,url,"
                        "language" if k == "repos" else
                        "repository,path,url" if k == "code" else
                        "title,url,state,number"])
        if out:
            try:
                data = json.loads(out)
                return _fmt_gh_search(data, k)
            except Exception:
                return out[:6000]
        return _http_github_search(query, k, limit, token, timeout, err)

    if not repo:
        return ("[错误] 需要 repo 参数，格式 owner/name，"
                "例如 lhclhc123/AIWorkbench")

    if act == "repo":
        out, err = _gh(["api", f"repos/{repo}"])
        if out:
            return _fmt_gh_repo(json.loads(out))
        return _http_github(f"https://api.github.com/repos/{repo}", token,
                            timeout, err, _fmt_gh_repo)

    if act in ("read", "list"):
        out, err = _gh(["api", f"repos/{repo}/contents/{path or ''}"
                        + (f"?ref={ref}" if ref else "")])
        if out:
            try:
                data = json.loads(out)
            except Exception:
                data = None
            if data is not None:
                return _fmt_gh_contents(repo, data, path, ref)
        return _http_github_contents(repo, path, ref, token, timeout, err)

    if act in ("issues", "pr"):
        want_pr = (act == "pr")
        q = f"repos/{repo}/issues?state={state}&per_page={int(limit or 20)}"
        out, err = _gh(["api", q])
        if out:
            try:
                items = json.loads(out)
            except Exception:
                items = []
            items = [x for x in items if bool(x.get("pull_request")) == want_pr]
            return _fmt_gh_issues(items, repo, state, want_pr)
        return f"[错误] 取 issue 失败：{err}"

    if act == "releases":
        out, err = _gh(["release", "list", "-R", repo, "-L", str(int(limit or 10))])
        if out:
            return f"[成功] {repo} 的 Release：\n" + out.strip()[:4000]
        return f"[错误] 取 Release 失败：{err}"

    return (f"[错误] 不认识的 action：{action}。可用：read / list / repo / "
            f"search / issues / pr / releases")


def _fmt_gh_repo(j):
    return ("[成功] 仓库信息\n"
            f"名称：{j.get('full_name')}\n"
            f"说明：{j.get('description')}\n"
            f"语言：{j.get('language')}  星标：{j.get('stargazers_count')}  "
            f"Fork：{j.get('forks_count')}\n"
            f"默认分支：{j.get('default_branch')}  "
            f"私有：{j.get('private')}\n"
            f"最近更新：{j.get('updated_at')}\n"
            f"地址：{j.get('html_url')}")


def _fmt_gh_search(data, kind):
    if not data:
        return "[成功] 没有搜到结果"
    lines = [f"[成功] 搜索到 {len(data)} 条（{kind}）："]
    for it in data[:20]:
        if kind == "repos":
            lines.append(f"· {it.get('fullName')}  ⭐{it.get('stargazersCount')}  "
                         f"[{it.get('language') or '?'}]  {it.get('url')}")
            if it.get("description"):
                lines.append(f"    {it['description'][:120]}")
        elif kind == "code":
            lines.append(f"· {it.get('repository', {}).get('nameWithOwner')} :: "
                         f"{it.get('path')}")
            lines.append(f"    {it.get('url')}")
        else:
            lines.append(f"· #{it.get('number')} [{it.get('state')}] "
                         f"{it.get('title')}")
            lines.append(f"    {it.get('url')}")
    return "\n".join(lines)


def _fmt_gh_contents(repo, data, path, ref):
    if isinstance(data, list):
        lines = [f"[成功] {repo}/{(path or '')}（{ref or '默认分支'}）共 "
                 f"{len(data)} 项："]
        for it in data:
            icon = "📁" if it.get("type") == "dir" else "📄"
            size = it.get("size") or 0
            lines.append(f"{icon} {it.get('name')}"
                         + (f"  （{size} B）" if it.get("type") == "file" else ""))
        return "\n".join(lines)
    # 单文件
    import base64
    content = data.get("content") or ""
    try:
        txt = base64.b64decode(content).decode("utf-8", "replace")
    except Exception:
        txt = ""
    if len(txt) > 12000:
        txt = txt[:12000] + f"\n…（共 {len(txt)} 字符，已截断）"
    return (f"[成功] {repo}/{data.get('path')}（{data.get('size')} B，"
            f"{ref or '默认分支'}）\n--- 内容 ---\n{txt}")


def _fmt_gh_issues(items, repo, state, want_pr):
    kind = "Pull Request" if want_pr else "Issue"
    if not items:
        return f"[成功] {repo} 没有 {state} 状态的 {kind}"
    lines = [f"[成功] {repo} 的 {kind}（状态 {state}）共 {len(items)} 条："]
    for it in items[:25]:
        lines.append(f"· #{it.get('number')} [{it.get('state')}] "
                     f"{it.get('title')}")
        if it.get("user"):
            lines.append(f"    作者：{it['user'].get('login')}  "
                         f"更新：{it.get('updated_at')}")
    return "\n".join(lines)


def _http_github(url, token, timeout, prev_err, fmt):
    """无 gh 时的公开 API 兜底。"""
    try:
        import requests
        h = {"Accept": "application/vnd.github+json",
             "User-Agent": "AIWorkbench"}
        if token:
            h["Authorization"] = f"Bearer {token}"
        r = requests.get(url, headers=h, timeout=int(timeout or 45))
        if r.status_code >= 400:
            return (f"[错误] GitHub 接口返回 {r.status_code}。"
                    f"（gh CLI 也不可用：{str(prev_err)[:120]}）"
                    f"\n提示：私有仓库 / 搜索 需要 GitHub Token 或登录 gh。")
        return fmt(r.json())
    except Exception as e:
        return (f"[错误] 访问 GitHub 失败：{type(e).__name__}: {e}\n"
                f"（gh CLI 也不可用：{str(prev_err)[:120]}）")


def _http_github_contents(repo, path, ref, token, timeout, prev_err):
    # 先用 raw 拿文件（公开仓库免认证）
    if path:
        base = f"https://raw.githubusercontent.com/{repo}/{ref or 'HEAD'}/{path}"
        try:
            import requests
            r = requests.get(base, timeout=int(timeout or 45),
                             headers={"User-Agent": "AIWorkbench"})
            if r.status_code == 200:
                txt = r.content.decode("utf-8", "replace")
                if len(txt) > 12000:
                    txt = txt[:12000] + f"\n…（共 {len(txt)} 字符，已截断）"
                return f"[成功] {repo}/{path}（raw，{len(r.content)} B）\n--- 内容 ---\n{txt}"
        except Exception:
            pass
    return _http_github(
        f"https://api.github.com/repos/{repo}/contents/{path or ''}"
        + (f"?ref={ref}" if ref else ""),
        token, timeout, prev_err,
        lambda j: _fmt_gh_contents(repo, j, path, ref))


def _http_github_search(query, kind, limit, token, timeout, prev_err):
    return _http_github(
        f"https://api.github.com/search/{kind}?q="
        + _urlq(str(query)) + f"&per_page={int(limit or 20)}",
        token, timeout, prev_err,
        lambda j: _fmt_gh_search(j.get("items") or [], kind))


def _urlq(s):
    from urllib.parse import quote
    return quote(s, safe="")
